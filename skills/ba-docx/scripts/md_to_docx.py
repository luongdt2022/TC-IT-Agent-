#!/usr/bin/env python3
"""Convert a practical Markdown SRS into a reviewable DOCX.

This is intentionally conservative: it converts common Markdown structures
without rewriting document content. Markdown remains the source of truth.
"""
from __future__ import annotations

import argparse
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(re.sub(r"`([^`]*)`", r"\1", text.strip()))
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if i == start + 1 and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            value = row[ci] if ci < len(row) else ""
            set_cell_text(table.cell(ri, ci), value, bold=ri == 0, color=(255, 255, 255) if ri == 0 else None)
            if ri == 0:
                set_cell_shading(table.cell(ri, ci), "1F4E78")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_inline_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    # Keep Markdown emphasis readable without attempting full CommonMark parsing.
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)")
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = p.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        else:
            p.add_run(token[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


def convert(src: Path, out: Path):
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    sec.left_margin = sec.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size, color in [("Title", 18, "1F4E78"), ("Heading 1", 15, "1F4E78"), ("Heading 2", 12, "2F75B5"), ("Heading 3", 10.5, "5B9BD5")]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    footer = sec.footer.paragraphs[0]
    footer.text = f"Nguồn: {src.name} · Bản Word phục vụ rà soát"
    add_page_number(footer)

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            level = min(len(h.group(1)), 3)
            style = "Title" if level == 1 else f"Heading {level}"
            add_inline_paragraph(doc, h.group(2).strip(), style=style)
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_paragraph(doc, re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
            # remove the empty paragraph created above
            p._element.getparent().remove(p._element)
            i += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            add_inline_paragraph(doc, re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
            i += 1
            continue
        if line.strip().startswith(">"):
            p = add_inline_paragraph(doc, line.strip()[1:].strip())
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            i += 1
            continue
        add_inline_paragraph(doc, line)
        i += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main():
    ap = argparse.ArgumentParser(description="Convert Markdown to review-ready DOCX")
    ap.add_argument("source", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    if not args.source.exists():
        ap.error(f"Markdown source not found: {args.source}")
    convert(args.source, args.output)
    print(f"created: {args.output}")


if __name__ == "__main__":
    main()
